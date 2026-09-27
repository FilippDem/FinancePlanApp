"""Plan reports: PDF (reportlab + matplotlib), Excel, CSV and JSON.

Sections are chosen by the user (report_data.SECTIONS). The PDF is landscape so the
year-by-year tables fit; money is in today's dollars unless `today=False`.
"""
from __future__ import annotations

import csv
import io
import json
from datetime import date

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.ticker import FuncFormatter  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import letter, landscape  # noqa: E402
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle  # noqa: E402
from reportlab.lib.units import inch  # noqa: E402
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak,  # noqa: E402
                                KeepTogether, CondPageBreak)

from server import report_data as RD  # noqa: E402

S = ['#2a78d6', '#eb6834', '#1baf7a', '#eda100', '#e87ba4', '#008300', '#4a3aa7', '#e34948']
INK, MUTED, GRID, ZEBRA = '#111318', '#80838c', '#e8e8e5', '#f7f7f5'
PAGE = landscape(letter)
W = PAGE[0] - 1.1 * inch


def _m(v, full=False):
    if v is None or v == '':
        return '—'
    if not isinstance(v, (int, float)):
        return str(v)
    if full:
        return f"{'-' if v < 0 else ''}${abs(v):,.0f}"
    a = abs(v); s = '-' if v < 0 else ''
    if a >= 1e6:
        return f"{s}${a / 1e6:.2f}M"
    if a >= 1e4:
        return f"{s}${a / 1e3:.0f}k"
    if a < 0.5:
        return '—'
    return f"{s}${a:,.0f}"


def _axis(ax):
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: _m(v)))
    for sp in ('top', 'right', 'left'):
        ax.spines[sp].set_visible(False)
    ax.spines['bottom'].set_color(GRID)
    ax.grid(axis='y', color=GRID, linewidth=0.8)
    ax.tick_params(colors=MUTED, labelsize=8, length=0)


def _img(fig, w, h):
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=150, bbox_inches='tight')
    plt.close(fig)
    buf.seek(0)
    return Image(buf, width=w * inch, height=h * inch)


class _Styles:
    def __init__(self):
        ss = getSampleStyleSheet()
        self.H1 = ParagraphStyle('h1', parent=ss['Title'], fontName='Helvetica-Bold', fontSize=22, textColor=colors.HexColor(INK), alignment=0, spaceAfter=4)
        self.H2 = ParagraphStyle('h2', parent=ss['Heading2'], fontName='Helvetica-Bold', fontSize=13, textColor=colors.HexColor(INK), spaceBefore=10, spaceAfter=5)
        self.H3 = ParagraphStyle('h3', parent=ss['Heading3'], fontName='Helvetica-Bold', fontSize=10.5, textColor=colors.HexColor(INK), spaceBefore=6, spaceAfter=3)
        self.B = ParagraphStyle('b', parent=ss['BodyText'], fontSize=9.5, leading=13, textColor=colors.HexColor('#333333'))
        self.Mu = ParagraphStyle('mu', parent=self.B, textColor=colors.HexColor(MUTED), fontSize=8.5, leading=11)


def _table(data, widths=None, font=8, right_from=1, header=True, zebra=True, group_rows=()):
    t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    st = [('FONTSIZE', (0, 0), (-1, -1), font), ('TOPPADDING', (0, 0), (-1, -1), 1.6), ('BOTTOMPADDING', (0, 0), (-1, -1), 1.6),
          ('LEFTPADDING', (0, 0), (-1, -1), 3), ('RIGHTPADDING', (0, 0), (-1, -1), 3),
          ('ALIGN', (right_from, 0), (-1, -1), 'RIGHT'), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]
    if header:
        st += [('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor(MUTED)), ('LINEBELOW', (0, 0), (-1, 0), 0.6, colors.HexColor(GRID)),
               ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold')]
    start = 1 if header else 0
    if zebra:
        for i in range(start, len(data)):
            if (i - start) % 2 == 1:
                st.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor(ZEBRA)))
    for i in group_rows:
        st += [('FONTNAME', (0, i), (-1, i), 'Helvetica-Bold'), ('BACKGROUND', (0, i), (-1, i), colors.HexColor('#eef3fb')),
               ('SPAN', (0, i), (-1, i)), ('ALIGN', (0, i), (-1, i), 'LEFT')]
    t.setStyle(TableStyle(st))
    return t


def _kv_table(pairs, st, cols=2):
    rows, cur = [], []
    for k, v in pairs:
        cur += [Paragraph(f"<font color='{MUTED}'>{k}</font>", st.B), Paragraph(str(v), st.B)]
        if len(cur) == 2 * cols:
            rows.append(cur); cur = []
    if cur:
        rows.append(cur + [''] * (2 * cols - len(cur)))
    t = Table(rows, colWidths=[W / (2 * cols) * 0.9, W / (2 * cols) * 1.1] * cols)
    t.setStyle(TableStyle([('LINEBELOW', (0, 0), (-1, -1), 0.4, colors.HexColor(GRID)), ('TOPPADDING', (0, 0), (-1, -1), 2),
                           ('BOTTOMPADDING', (0, 0), (-1, -1), 2)]))
    return t


def _charts(ctx, st, today):
    p, pr, mc = ctx['plan'], ctx['proj'], ctx['mc']
    rows, s = pr['rows'], pr['summary']
    idx = [r['infl_index'] if today else 1.0 for r in rows]
    years = [r['year'] for r in rows]
    single = ctx['single']
    out = []
    fig, ax = plt.subplots(figsize=(11, 3.4))
    band = lambda q: [v / i for v, i in zip(mc['net_worth'][q], idx)]
    exp_path = [r['net_worth'] / f for r, f in zip(rows, idx)]
    ax.fill_between(years, band('10'), band('90'), color=S[0], alpha=0.10, linewidth=0, label='10th–90th percentile')
    ax.fill_between(years, band('25'), band('75'), color=S[0], alpha=0.20, linewidth=0, label='25th–75th percentile')
    ax.plot(years, band('50'), color=S[0], linewidth=2, label='Median')
    ax.plot(years, exp_path, color=S[1], linewidth=1.6, label='Expected path')
    ax.axhline(0, color=MUTED, linewidth=0.8)
    top = max(max(band('75')), max(exp_path), 1.0) * 1.25
    bot = min(min(band('25')), min(exp_path), 0.0)
    ax.set_ylim(bot - 0.05 * top if bot < 0 else 0, top)
    marks = [(s['retirement_year1'], ctx['names'][0])] + ([] if single else [(s['retirement_year2'], ctx['names'][1])])
    if len(marks) == 2 and abs(marks[0][0] - marks[1][0]) <= 4:
        marks = [(min(m[0] for m in marks), 'Both retire')]
    for y, who in marks:
        ax.axvline(y, color=MUTED, linewidth=0.8, alpha=0.6)
        ax.text(y + 0.5, top * 0.97, who if who == 'Both retire' else f"{who} retires", fontsize=7, color=MUTED, va='top')
    _axis(ax)
    ax.legend(frameon=False, fontsize=7.5, ncol=4, loc='upper left', bbox_to_anchor=(0, 1.13))
    out += [KeepTogether([Paragraph('Net worth', st.H2), _img(fig, 9.6, 3.0)])]

    fig, ax = plt.subplots(figsize=(11, 3.2))
    cats = [('Living', lambda r: r['exp_person1'] + r['exp_person2'] + r['exp_family']), ('Housing', lambda r: r['exp_housing']),
            ('Children', lambda r: r['exp_children']), ('Healthcare', lambda r: r['exp_healthcare']),
            ('Recurring & one-time', lambda r: r['exp_recurring'] + r['exp_purchases'] + r['down_payment']), ('Taxes', lambda r: r['taxes'])]
    bottom = [0.0] * len(rows)
    for k, (lab, f) in enumerate(cats):
        vals = [f(r) / i for r, i in zip(rows, idx)]
        ax.bar(years, vals, bottom=bottom, color=S[k], width=0.8, label=lab, edgecolor='white', linewidth=0.4)
        bottom = [b + v for b, v in zip(bottom, vals)]
    ax.step(years, [r['total_income'] / i for r, i in zip(rows, idx)], where='mid', color=INK, linewidth=1.6, label='Income')
    _axis(ax)
    ax.legend(frameon=False, fontsize=7.5, ncol=7, loc='upper left', bbox_to_anchor=(0, 1.13))
    out += [KeepTogether([Paragraph('Income and spending', st.H2), _img(fig, 9.6, 2.8)])]
    return out


class _Cites:
    """Numbers sources in order of first use; markers link to the Sources list at the end."""
    def __init__(self, ctx, sections):
        self.ctx, self.order = ctx, [i for i, _ in RD.sources_for(ctx, sections)]
        self.data = dict(RD.sources_for(ctx, sections))

    def __call__(self, section: str) -> str:
        ids = [i for i in RD.section_sources(self.ctx, section) if i in self.data]
        if not ids:
            return ''
        return '<super>' + ''.join(f'<a href="#src_{i}" color="{S[0]}">[{self.order.index(i) + 1}]</a>' for i in ids) + '</super>'

    def story(self, st) -> list:
        if not self.order:
            return []
        from xml.sax.saxutils import escape as x
        out = [CondPageBreak(1.5 * inch), Paragraph('Sources', st.H2)]
        for k, i in enumerate(self.order, 1):
            d = self.data[i]
            note = f" <i>{x(d['note'])}.</i>" if d.get('note') else ''
            out.append(Paragraph(f'<a name="src_{i}"/>[{k}] {x(d["publisher"])}, <link href="{x(d["url"])}" color="{S[0]}">{x(d["title"])}</link> '
                                 f'({x(d["year"])}). Used for: {x(d.get("used_for", ""))}.{note} '
                                 f'<font color="{MUTED}">{x(d["url"])}</font>', st.Mu))
        return out


def build(plan: dict, household: str, checkins: list | None = None, sections: list | None = None,
          title: str | None = None, today: bool = True, ctx: dict | None = None) -> bytes:
    sections = [s for s in (sections or RD.DEFAULT_SECTIONS) if s in RD.SECTIONS]
    ctx = ctx or RD.compute(plan)
    p, pr, mc = ctx['plan'], ctx['proj'], ctx['mc']
    rows, s = pr['rows'], pr['summary']
    single, (n1, n2) = ctx['single'], ctx['names']
    st = _Styles()
    cite = _Cites(ctx, sections)
    dollars = "today's dollars" if today else 'nominal (future) dollars'
    names = n1 + ('' if single else f" & {n2}")
    story = [Paragraph(title or f"Financial plan — {household}", st.H1),
             Paragraph(f"{names} · prepared {date.today():%B %d, %Y} · amounts in {dollars} unless noted", st.Mu), Spacer(1, 8)]

    if 'summary' in sections:
        ret = next((r for r in rows if r['year'] == s['retirement_year']), rows[-1])
        rf = ret['infl_index'] if today else 1.0
        kpis = [('Net worth today', _m(s['net_worth_now'])), (f"At retirement ({s['retirement_year']})", _m(ret['net_worth'] / rf)),
                ('Plan success', f"{mc['success_rate'] * 100:.0f}%"),
                ('Expected path: savings last', f"until {s['depletion_year']}" if s['depletion_year'] else 'for life'),
                ('Income this year', _m(rows[0]['total_income'])), ('Spending this year', _m(rows[0]['total_expenses']))]
        t = Table([[Paragraph(f"<font size=8 color='{MUTED}'>{k}</font><br/><font size=14><b>{v}</b></font>", st.B) for k, v in kpis]],
                  colWidths=[W / 6] * 6)
        t.setStyle(TableStyle([('BOX', (0, 0), (-1, -1), 0.6, colors.HexColor(GRID)), ('INNERGRID', (0, 0), (-1, -1), 0.6, colors.HexColor(GRID)),
                               ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7)]))
        story += [t, Paragraph('Lifetime summary', st.H2), _kv_table(RD.lifetime_summary(ctx, today), st, cols=3)]

    if 'charts' in sections:
        story += _charts(ctx, st, today)

    if 'people' in sections:
        hdr, data = RD.people_table(ctx)
        fmt = [[r[0], r[1], _m(r[2], True), f"{r[3]}%", _m(r[4], True), _m(r[5], True), r[6], _m(r[7], True), r[8] or 'at retirement', r[9]] for r in data]
        story += [CondPageBreak(2.5 * inch), Paragraph('People & income' + cite('people'), st.H2), _table([hdr] + fmt, right_from=1)]
        cr = RD.career_rows(ctx)
        if cr:
            story += [Paragraph('Career phases & job changes', st.H3),
                      _table([['Person', 'Phase', 'Ages / year', 'Salary', 'Raise %', 'Bonus %', 'RSU /yr']] +
                             [[a, b, c, _m(d, True), e, f, _m(g, True) if g != '' else ''] for a, b, c, d, e, f, g in cr])]

    if 'children' in sections and p['children_list']:
        hdr, data = RD.children_table(ctx)
        story += [CondPageBreak(1.5 * inch), Paragraph('Children' + cite('children'), st.H2), _table([hdr] + data, right_from=1)]

    if 'spending' in sections:
        items = RD.spending_inputs(ctx)
        story += [CondPageBreak(2 * inch), Paragraph("Spending plan (per year, today's dollars, before moves and inflation)" + cite('spending'), st.H2)]
        groups = {}
        for g, k, v in items:
            groups.setdefault(g, []).append((k, v))
        data, grp_rows = [['Category', 'Per year', 'Per month']], []
        for g, lst in groups.items():
            grp_rows.append(len(data)); data.append([f"{g} — {_m(sum(v for _, v in lst), True)}/yr", '', ''])
            data += [[k, _m(v, True), _m(v / 12, True)] for k, v in lst]
        story.append(_table(data, widths=[3.2 * inch, 1.3 * inch, 1.3 * inch], group_rows=grp_rows, zebra=False))

    if 'homes' in sections and p['houses']:
        hdr, data = RD.homes_table(ctx)
        fmt = [[r[0], r[1], _m(r[2], True), _m(r[3], True), _m(r[4], True), r[5], r[6], r[7], _m(r[8], True), Paragraph(r[9], st.Mu)] for r in data]
        story += [CondPageBreak(1.5 * inch), Paragraph('Homes & mortgages', st.H2),
                  _table([hdr] + fmt, widths=[1.5 * inch] + [0.75 * inch] * 8 + [2.3 * inch])]

    if 'healthcare' in sections:
        hc = RD.healthcare_tables(ctx)
        story += [CondPageBreak(1.5 * inch), Paragraph('Healthcare & insurance' + cite('healthcare'), st.H2)]
        if hc['insurance'][1]:
            h, d = hc['insurance']
            story.append(_table([h] + [[r[0], r[1], _m(r[2], True), _m(r[3], True), _m(r[4], True), r[5], r[6]] for r in d]))
        if hc['ltc'][1]:
            h, d = hc['ltc']
            story += [Paragraph('Long-term care', st.H3), _table([h] + [[r[0], r[1], _m(r[2], True), _m(r[3], True), r[4], r[5]] for r in d])]
        story += [Spacer(1, 4), _kv_table([(k, _m(v, True)) for k, v in hc['other']], st, cols=3)]

    if 'purchases' in sections:
        pt = RD.purchases_tables(ctx)
        if pt['purchases'][1] or pt['recurring'][1]:
            story += [CondPageBreak(1.5 * inch), Paragraph('One-time purchases & recurring costs', st.H2)]
        if pt['purchases'][1]:
            h, d = pt['purchases']
            story.append(_table([h] + [[r[0], r[1], _m(r[2], True), r[3], r[4], r[5]] for r in d]))
        if pt['recurring'][1]:
            h, d = pt['recurring']
            story += [Spacer(1, 4), _table([h] + [[r[0], r[1], _m(r[2], True), r[3], r[4], r[5], r[6]] for r in d])]

    if 'locations' in sections:
        h, d = RD.locations_table(ctx)
        story += [CondPageBreak(1.2 * inch), Paragraph('Where you live' + cite('locations'), st.H2), _table([h] + d)]

    if 'assumptions' in sections:
        story += [CondPageBreak(1.2 * inch), Paragraph('Key assumptions' + cite('assumptions'), st.H2), _kv_table(RD.assumptions_rows(ctx), st, cols=3)]

    if 'monte_carlo' in sections:
        h, d = RD.mc_percentiles(ctx, today)
        step = 5 if len(d) > 30 else 1
        sel = [r for i, r in enumerate(d) if i % step == 0 or i == len(d) - 1]
        story += [CondPageBreak(2.5 * inch), Paragraph('Monte Carlo results' + cite('monte_carlo'), st.H2),
                  Paragraph(f"{mc['n']:,} simulations · success rate {mc['success_rate'] * 100:.0f}% · final net worth median "
                            f"{_m(mc['final']['median'] / (rows[-1]['infl_index'] if today else 1))}, 10th percentile {_m(mc['final']['p10'] / (rows[-1]['infl_index'] if today else 1))} "
                            f"({'today' if today else 'nominal'}'s dollars){' · every 5th year shown' if step > 1 else ''}".replace("nominal's", 'nominal'), st.Mu),
                  _table([h] + [[r[0]] + [_m(v) for v in r[1:-1]] + [f"{r[-1] * 100:.0f}%"] for r in sel])]

    if 'year_by_year' in sections:
        hdr, data = RD.category_table(ctx, today)
        keys = [k for k, _, _ in RD.CATEGORY_COLS if not (single and k == 'p2')]
        pick = lambda wanted: [keys.index(k) for k in wanted if k in keys]
        a_cols = pick(['year', 'ages', 'wages', 'ss', 'rent', 'sales', 'growth', 'taxes', 'spend', 'cashflow', 'savings', 'equity', 'nw'])
        b_cols = pick(['year', 'ages', 'p1', 'p2', 'family', 'children', 'housing', 'healthcare', 'recurring', 'onetime', 'down', 'spend'])
        fmt = lambda r, cols: [r[i] if keys[i] in ('year', 'ages') else _m(r[i]) for i in cols]
        neg = lambda cols: [(c, i + 1) for i, r in enumerate(data) for c, ci in enumerate(cols)
                            if isinstance(r[ci], (int, float)) and keys[ci] in ('cashflow', 'savings', 'nw') and r[ci] < 0]
        for title_, cols in ((f'Year by year: income, taxes and balances ({dollars})', a_cols),
                             (f'Year by year: spending by category ({dollars})', b_cols)):
            t = _table([[hdr[i] for i in cols]] + [fmt(r, cols) for r in data], widths=[W / len(cols)] * len(cols), font=7, right_from=2)
            t.setStyle(TableStyle([('TEXTCOLOR', (c, i), (c, i), colors.HexColor('#d03b3b')) for c, i in neg(cols)]))
            story += [PageBreak(), Paragraph(title_, st.H2), t]

    if 'category_detail' in sections:
        years, lines = RD.detail_lines(ctx, today)
        story += [PageBreak(), Paragraph(f'Detailed breakdown: every line item ({dollars})', st.H2),
                  Paragraph('Each block covers ten years. Lines that are zero for the whole plan are left out.', st.Mu)]
        for start in range(0, len(years), 10):
            yrs = years[start:start + 10]
            data, grp_rows, last_g = [['Line item'] + [str(y) for y in yrs]], [], None
            for g, lab, vals in lines:
                chunk = vals[start:start + 10]
                if not any(abs(v) >= 0.5 for v in chunk):
                    continue
                if g != last_g:
                    grp_rows.append(len(data)); data.append([g] + [''] * len(yrs)); last_g = g
                data.append([lab] + [_m(v) for v in chunk])
            widths = [2.6 * inch] + [(W - 2.6 * inch) / len(yrs)] * len(yrs)
            story += [CondPageBreak(3 * inch), Paragraph(f"{yrs[0]}–{yrs[-1]}", st.H3),
                      _table(data, widths=widths, font=6.8, group_rows=grp_rows, zebra=False)]

    if 'checkins' in sections and checkins:
        cd = [['Date', 'Type', 'Savings', 'Expected', 'Net worth', 'Status', 'Percentile']]
        for c in checkins:
            cd.append([c.get('date'), c.get('kind'), _m((c.get('totals') or {}).get('investable')), _m((c.get('expected') or {}).get('investable')),
                       _m((c.get('totals') or {}).get('net_worth')), (c.get('status') or '').replace('_', ' '),
                       f"{c['percentile']:.0f}" if isinstance(c.get('percentile'), (int, float)) else ''])
        story += [CondPageBreak(1.5 * inch), Paragraph('Check-in history', st.H2), _table(cd)]

    story += cite.story(st)
    buf = io.BytesIO()

    def footer(canvas, doc):
        canvas.setFont('Helvetica', 7.5)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(0.55 * inch, 0.4 * inch, f"Financial Planning Suite · {household} · projections are estimates, not advice")
        canvas.drawRightString(PAGE[0] - 0.55 * inch, 0.4 * inch, f"Page {doc.page}")

    doc = SimpleDocTemplate(buf, pagesize=PAGE, leftMargin=0.55 * inch, rightMargin=0.55 * inch, topMargin=0.5 * inch,
                            bottomMargin=0.65 * inch, title=title or f"Financial plan — {household}")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ── other formats ───────────────────────────────────────────────────────
def build_csv(plan: dict, today: bool = True, ctx: dict | None = None, detail: bool = False) -> bytes:
    ctx = ctx or RD.compute(plan, n_mc=200)
    buf = io.StringIO()
    w = csv.writer(buf)
    if detail:
        years, lines = RD.detail_lines(ctx, today)
        w.writerow(['Group', 'Line item'] + years)
        for g, lab, vals in lines:
            w.writerow([g, lab] + [round(v, 2) for v in vals])
    else:
        hdr, data = RD.category_table(ctx, today)
        w.writerow(hdr)
        for r in data:
            w.writerow([round(v, 2) if isinstance(v, float) else v for v in r])
    return buf.getvalue().encode('utf-8-sig')


def build_json(plan: dict, today: bool = True, ctx: dict | None = None, checkins: list | None = None) -> bytes:
    ctx = ctx or RD.compute(plan)
    hdr, data = RD.category_table(ctx, today)
    years, lines = RD.detail_lines(ctx, today)
    mh, md = RD.mc_percentiles(ctx, today)
    out = {
        'generated': date.today().isoformat(), 'dollars': 'today' if today else 'nominal',
        'plan': ctx['plan'], 'summary': ctx['proj']['summary'], 'lifetime': dict(RD.lifetime_summary(ctx, today)),
        'year_by_year': [dict(zip(hdr, r)) for r in data],
        'line_items': {'years': years, 'lines': [{'group': g, 'item': l, 'values': v} for g, l, v in lines]},
        'monte_carlo': {'n': ctx['mc']['n'], 'success_rate': ctx['mc']['success_rate'], 'final': ctx['mc']['final'],
                        'percentiles_by_year': [dict(zip(mh, r)) for r in md]},
        'events': ctx['proj']['events'], 'checkins': checkins or [],
        'sources': [{'n': k, 'id': i, **d} for k, (i, d) in enumerate(RD.sources_for(ctx, list(RD.SECTIONS)), 1)],
    }
    return json.dumps(out, indent=1, default=str).encode()


def build_xlsx(plan: dict, household: str, today: bool = True, ctx: dict | None = None, checkins: list | None = None,
               sections: list | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    sections = sections or RD.DEFAULT_SECTIONS
    ctx = ctx or RD.compute(plan)
    HDR = PatternFill(start_color='2A78D6', end_color='2A78D6', fill_type='solid')
    GRP = PatternFill(start_color='EEF3FB', end_color='EEF3FB', fill_type='solid')
    CUR = '"$"#,##0;[Red]-"$"#,##0'
    wb = Workbook()

    import re as _re
    _NUM = _re.compile(r'\b(age|ages|age now|year|years|born|bought|retire at|claim at|plan until|yrs|days|from|to|every)\b')

    def sheet(title, header, rows, money_from=1, widths=None, first=False, int_money=True):
        ws = wb.active if first else wb.create_sheet(title[:31])
        ws.title = title[:31]
        ws.append(header)
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font, cell.alignment = HDR, Font(bold=True, color='FFFFFF'), Alignment(horizontal='center', wrap_text=True)
        for r in rows:
            ws.append(list(r))
        # pick a number format from the column header: percents, ages/years/counts, else currency
        fmts = []
        for h in header:
            hl = str(h).lower()
            if 'solvent' in hl:
                fmts.append('0%')
            elif '%' in hl or 'rate' in hl:
                fmts.append('0.0"%"')
            elif _NUM.search(hl) and '$' not in hl:
                fmts.append('0')
            else:
                fmts.append(CUR)
        for row in ws.iter_rows(min_row=2):
            for i, c in enumerate(row):
                if i < money_from or isinstance(c.value, bool) or not isinstance(c.value, (int, float)):
                    continue
                if isinstance(c.value, float) or fmts[i] != CUR or int_money:
                    c.number_format = fmts[i] if i < len(fmts) else CUR
        for i in range(1, len(header) + 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 14)
        ws.freeze_panes = 'B2'
        return ws

    s = ctx['proj']['summary']
    summary_rows = [('Household', household), ('Prepared', date.today().isoformat()), ('Dollars', "today's" if today else 'nominal'),
                    ('Plan success (Monte Carlo)', f"{ctx['mc']['success_rate'] * 100:.0f}%"),
                    ('Net worth today', s['net_worth_now']), ('Retirement year', s['retirement_year'])] + RD.lifetime_summary(ctx, today)
    sheet('Summary', ['Item', 'Value'], summary_rows, widths={1: 34, 2: 30}, first=True, int_money=False)
    hdr, data = RD.category_table(ctx, today)
    sheet('Year by year', hdr, data, money_from=2)
    years, lines = RD.detail_lines(ctx, today)
    ws = sheet('Line items', ['Group', 'Line item'] + years, [[g, l] + v for g, l, v in lines], money_from=2, widths={1: 22, 2: 34})
    if 'people' in sections:
        h, d = RD.people_table(ctx)
        sheet('People', h, d, money_from=1)
    if 'children' in sections and ctx['plan']['children_list']:
        h, d = RD.children_table(ctx)
        sheet('Children', h, d)
    if 'spending' in sections:
        sheet('Spending plan', ['Who', 'Category', 'Per year'], RD.spending_inputs(ctx), money_from=2, widths={1: 18, 2: 28})
    if 'homes' in sections and ctx['plan']['houses']:
        h, d = RD.homes_table(ctx)
        sheet('Homes', h, d, widths={10: 60})
    if 'healthcare' in sections:
        hc = RD.healthcare_tables(ctx)
        sheet('Health insurance', *hc['insurance'])
        if hc['ltc'][1]:
            sheet('Long-term care', *hc['ltc'])
    if 'purchases' in sections:
        pt = RD.purchases_tables(ctx)
        sheet('One-time purchases', *pt['purchases'])
        sheet('Recurring', *pt['recurring'])
    if 'locations' in sections:
        sheet('Locations', *RD.locations_table(ctx))
    if 'assumptions' in sections:
        sheet('Assumptions', ['Assumption', 'Value'], RD.assumptions_rows(ctx), widths={1: 30, 2: 30})
    if 'monte_carlo' in sections:
        h, d = RD.mc_percentiles(ctx, today)
        sheet('Monte Carlo', h, d)
    if 'checkins' in sections and checkins:
        sheet('Check-ins', ['Date', 'Type', 'Savings', 'Expected', 'Net worth', 'Status'],
              [[c.get('date'), c.get('kind'), (c.get('totals') or {}).get('investable'), (c.get('expected') or {}).get('investable'),
                (c.get('totals') or {}).get('net_worth'), c.get('status')] for c in checkins])
    srcs = RD.sources_for(ctx, sections)
    if srcs:
        ws = sheet('Sources', ['#', 'Publisher', 'Title', 'Year', 'Link', 'Used for'],
                   [[k, d['publisher'], d['title'], d['year'], d['url'], d.get('used_for', '') + (f" ({d['note']})" if d.get('note') else '')]
                    for k, (_, d) in enumerate(srcs, 1)], money_from=99, widths={2: 30, 3: 60, 4: 14, 5: 50, 6: 60})
        for r in range(2, len(srcs) + 2):
            c = ws.cell(row=r, column=5)
            c.hyperlink = c.value
            c.font = Font(color='2A78D6', underline='single')
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
