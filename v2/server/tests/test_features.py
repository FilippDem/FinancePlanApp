"""Endpoints added in the 'living plan' round: actuals, workbook, history, stress, retirement, report, email."""
import json
from datetime import date

from server.tests.test_api import client, v08_household  # noqa: F401  (fixture)


def _login(c, tmp):
    hid = v08_household(tmp)
    c.post('/api/login', json={'email': 'filippdem@gmail.com'})
    c.post('/api/households/select', json={'id': hid})
    return hid


def test_actuals_year_merge_keeps_other_fields(client):
    c, tmp = client
    hid = _login(c, tmp)
    r = c.put('/api/actuals/2025', json={'actual': {'net_worth': 500000, 'expenses': {'family': {'Food & Groceries': 9000}}}})
    assert r.status_code == 200
    c.put('/api/actuals/2025', json={'actual': {'expenses': {'family': {'Utilities': 3000}}}})
    act = c.get('/api/actuals').json()['actuals']['2025']
    assert act['x'] == 1 and act['net_worth'] == 500000
    assert act['expenses']['family'] == {'Food & Groceries': 9000, 'Utilities': 3000}
    stored = json.loads((tmp / 'households' / f'{hid}.json').read_text())
    assert stored['scenarios']['Base']['parent1_name'] == 'Filipp'


def test_planned_endpoint(client):
    c, tmp = client
    _login(c, tmp)
    plan = c.get('/api/plan').json()['plan']
    pl = c.post('/api/actuals/planned', json={'plan': plan, 'year': plan['current_year']}).json()
    assert pl['in_projection'] and pl['totals']['spending'] > 0 and 'family' in pl['group_totals']


def test_workbook_roundtrip(client):
    import io
    from openpyxl import load_workbook
    c, tmp = client
    _login(c, tmp)
    y = date.today().year
    r = c.get(f'/api/actuals/workbook.xlsx?start={y}&end={y}')
    assert r.status_code == 200 and r.content[:2] == b'PK'
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb[f'Expenses_{y}']
    # fill Jan+Feb for the first category row, and year-end net worth on Summary
    for row in ws.iter_rows(min_row=2):
        if row[0].value and not row[0].font.bold:
            row[2].value, row[3].value = 100, 150
            cat = row[0].value
            break
    wb['Summary']['C2'] = 777000
    buf = io.BytesIO(); wb.save(buf)
    imp = c.post('/api/actuals/import', content=buf.getvalue()).json()
    assert str(y) in imp['years']
    a = imp['actuals'][str(y)]
    assert a['net_worth'] == 777000
    assert a['expenses']['parentX'][cat] == 250
    assert imp['actuals']['2025'] == {'x': 1}   # untouched


def test_history_and_restore(client):
    c, tmp = client
    hid = _login(c, tmp)
    c.put('/api/plan', json={'plan': {'parentX_age': 40}})
    c.put('/api/plan', json={'plan': {'parentX_age': 41}})
    vs = c.get('/api/history').json()['versions']
    assert vs, 'backups should be listed'
    old = next(v for v in vs if v['kind'] == 'recent')
    prev = c.post('/api/history/preview', json={'id': old['id']}).json()['plan']
    r = c.post('/api/history/restore', json={'id': old['id']}).json()
    assert r['plan']['parentX_age'] == prev['parentX_age']
    assert r['plan']['future_v09_field'] == [1, 2, 3]
    stored = json.loads((tmp / 'households' / f'{hid}.json').read_text())
    assert 'restored' in stored['saved_by'] and stored['actuals'] == {'2025': {'x': 1}}
    assert c.post('/api/history/restore', json={'id': 'recent/../../etc'}).status_code == 400
    assert c.post('/api/history/restore', json={'id': 'recent/otherhh_2026.json'}).status_code == 400


def test_stress_and_retirement(client):
    c, tmp = client
    _login(c, tmp)
    plan = c.get('/api/plan').json()['plan']
    tests = c.post('/api/stress/defaults', json={'plan': plan}).json()['tests']
    assert {t['type'] for t in tests} >= {'market_crash', 'income_loss', 'extra_cost', 'inflation_spike', 'early_death'}
    r = c.post('/api/stress', json={'plan': plan, 'tests': tests[:2], 'n': 100}).json()
    assert len(r['results']) == 2 and 0 <= r['base']['success_rate'] <= 1
    rt = c.post('/api/retirement', json={'plan': plan}).json()
    assert rt['ss'] and 'ratio' in rt['replacement']
    wi = c.post('/api/retirement/whatif', json={'plan': plan, 'n': 100}).json()['rows']
    assert [w['delta'] for w in wi] == [-3, -2, -1, 0, 1, 2, 3]


def test_report_pdf(client):
    c, tmp = client
    _login(c, tmp)
    r = c.get('/api/report.pdf')
    assert r.status_code == 200 and r.content[:4] == b'%PDF'
    plan = c.get('/api/plan').json()['plan']
    plan['parentX_retirement_age'] = 60
    r = c.post('/api/report.pdf', json={'plan': plan})
    assert r.status_code == 200 and r.content[:4] == b'%PDF'
    secs = [x['key'] for x in c.get('/api/report/sections').json()['sections']]
    assert 'year_by_year' in secs and 'category_detail' in secs
    r = c.post('/api/report', json={'format': 'pdf', 'sections': ['summary', 'year_by_year'], 'title': 'My plan'})
    assert r.status_code == 200 and r.content[:4] == b'%PDF' and 'My-plan' in r.headers['content-disposition']
    x = c.post('/api/report', json={'format': 'xlsx'})
    assert x.content[:2] == b'PK'
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(x.content))
    assert {'Summary', 'Year by year', 'Line items', 'Monte Carlo'} <= set(wb.sheetnames)
    cs = c.post('/api/report', json={'format': 'csv'}).content.decode('utf-8-sig')
    assert cs.splitlines()[0].startswith('Year,Ages,Wages')
    cd = c.post('/api/report', json={'format': 'csv', 'detail': True}).content.decode('utf-8-sig')
    assert cd.startswith('Group,Line item')
    js = c.post('/api/report', json={'format': 'json'}).json()
    assert js['year_by_year'] and js['line_items']['lines'] and 'monte_carlo' in js


def test_email_settings_and_reminder_job(client, monkeypatch):
    c, tmp = client
    hid = _login(c, tmp)
    st = c.get('/api/notify/status').json()
    assert st['configured'] is False
    assert c.post('/api/checkins/test-email').status_code == 400
    s = c.put('/api/checkins/settings', json={'cadence': 'quarterly', 'email_reminders': True}).json()['settings']
    assert s['email_reminders'] is True
    # force due and run the job with a fake sender
    c.put('/api/checkins/settings', json={'next_due': '2026-01-01'})
    from server import notify
    sent = []
    today = date(2026, 1, 3)
    assert notify.run_once(today, sender=lambda to, subj, text, html: sent.append((to, subj))) == [hid]
    assert sent[0][0] == ['filippdem@gmail.com'] and 'due' in sent[0][1]
    # same due date: no second email until 7 days later, then exactly one follow-up
    assert notify.run_once(date(2026, 1, 5), sender=lambda *a: sent.append(a)) == []
    assert notify.run_once(date(2026, 1, 11), sender=lambda *a: sent.append(a)) == [hid]
    assert notify.run_once(date(2026, 1, 20), sender=lambda *a: sent.append(a)) == []
    assert len(sent) == 2
    # off → nothing
    c.put('/api/checkins/settings', json={'email_reminders': False, 'next_due': '2026-04-01'})
    assert notify.run_once(date(2026, 4, 3), sender=lambda *a: sent.append(a)) == []


def test_household_admin(client):
    c, tmp = client
    hid = _login(c, tmp)
    h = c.get('/api/household').json()
    assert h['id'] == hid and h['you'] == 'filippdem@gmail.com'
    assert c.put('/api/household', json={'name': 'Demenschonok'}).json()['name'] == 'Demenschonok'
    idx = json.loads((tmp / 'households_index.json').read_text())
    assert idx[hid]['name'] == 'Demenschonok'
    assert c.delete('/api/household/members/filippdem@gmail.com').status_code == 400   # last member
    demos = c.get('/api/demos').json()['demos']
    name = next(iter(demos))
    r = c.post('/api/demos/open', json={'name': name}).json()
    assert r['id'].startswith('_test_')
    assert c.get('/api/plan').json()['plan']['parent1_name'] == demos[name]['parent1_name']
    assert c.post('/api/households/cleanup-tests').json()['removed'] >= 1
    assert hid in json.loads((tmp / 'households_index.json').read_text())


def test_cited_sources_exist():
    """Every source id cited in the web app or reports is in sources.json with a link."""
    import re
    from pathlib import Path
    from finplan.reference import sources
    from server import report_data as RD
    src = sources()
    assert all(v.get('url', '').startswith('https://') and v.get('publisher') and v.get('title') for v in src.values())
    root = Path(__file__).resolve().parents[2] / 'web' / 'src'
    text = ''.join(p.read_text(encoding='utf-8') for p in root.rglob('*.tsx'))
    ids = set(re.findall(r"'([a-z0-9]+_[a-z0-9_]+)'", text)) | set(re.findall(r'id="([a-z0-9_]+)"', text))
    cited = {i for i in ids if i in src or re.match(r'^(bls|bea|cms|irs|ssa|kff|fed|oecd|hud|mit|worldbank|damodaran|taxfoundation|childcareaware|collegeboard|numbeo)', i)}
    assert cited and cited <= set(src), sorted(cited - set(src))
    ctx = {'plan': {'state_timeline': [{'state': 'Portugal'}], 'mc_use_historical': True}}
    for sec in RD.SECTIONS:
        assert set(RD.section_sources(ctx, sec)) <= set(src)
